"""
The study dossier (Word): what a methods section and its appendix need, built from the session's records.

Sections:

1. Provenance: the dataset, its source file and hash, the parse options, the data fingerprint, the trace and
   its integrity, the software versions;
2. Analysis population and participant flow (STROBE item 13): the flow diagram, the exclusions with the
   criteria behind them, and the definition of every criterion;
3. Data handling: the recorded steps that changed the data (merges, derived variables, conversions,
   clipping, imputations);
4. Missing data: missing values per variable, values filled by single imputation, and the multiple imputation
   with its methods paragraph (reporting items of Sterne et al. 2009);
5. Results: each result added to the dossier with what it was computed on;
6. References.

The input is plain data (:class:`DossierInput`); the app gathers it (app_pages/study_dossier.py). The dossier
supports these reporting guidelines; it does not certify compliance, and it is not an electronically signed record.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor

from core.attrition import FlowExit, ParticipantFlow
from manage.report_tables import ResultBlocks

METHOD_REFERENCES = {
    "strobe": "von Elm E, Altman DG, Egger M, Pocock SJ, Gotzsche PC, Vandenbroucke JP (2007). The Strengthening the "
    "Reporting of Observational Studies in Epidemiology (STROBE) statement: guidelines for reporting observational "
    "studies. Lancet 370:1453-1457.",
    "rubin1987": "Rubin DB (1987). Multiple Imputation for Nonresponse in Surveys. Wiley.",
    "barnard1999": "Barnard J, Rubin DB (1999). Small-sample degrees of freedom with multiple imputation. Biometrika "
    "86:948-955.",
    "vanbuuren2018": "van Buuren S (2018). Flexible Imputation of Missing Data, 2nd ed. Chapman & Hall/CRC.",
    "sterne2009": "Sterne JAC, White IR, Carlin JB, et al. (2009). Multiple imputation for missing data in "
    "epidemiological and clinical research: potential and pitfalls. BMJ 338:b2393.",
    "white2011": "White IR, Royston P, Wood AM (2011). Multiple imputation using chained equations: issues and "
    "guidance for practice. Stat Med 30:377-399.",
}


@dataclass
class DossierResult:
    title: str
    settings: str
    provenance: str
    blocks: ResultBlocks
    stale: str | None = (
        None  # why the result no longer describes the current data, when it was kept anyway
    )
    single_imputed: tuple[tuple[str, int], ...] = ()


@dataclass
class DossierInput:
    title: str
    flow: ParticipantFlow
    population_label: str
    author: str = ""
    generated_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="minutes"))
    provenance: list[tuple[str, str]] = field(default_factory=list)
    flow_png: bytes | None = None
    criteria: pd.DataFrame | None = None  # family, criterion, kind, definition
    data_steps: list[tuple[str, str]] = field(default_factory=list)  # (step, what it did)
    missing: pd.DataFrame | None = None  # variable, missing, share
    single_imputed: dict[str, int] = field(default_factory=dict)
    exact_imputed: dict[str, int] = field(default_factory=dict)  # values imputed by an exact formula
    mi: dict[str, Any] | None = None  # methods_paragraph, settings rows, methods rows, status
    results: list[DossierResult] = field(default_factory=list)
    references: list[str] = field(default_factory=list)
    omitted: list[str] = field(default_factory=list)  # stale results left out


def _small(paragraph: Any, text: str, *, italic: bool = True, size: float = 9) -> None:
    run = paragraph.add_run(text)
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(90, 90, 90)


def _table(doc: Any, frame: pd.DataFrame) -> None:
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(name)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = (
                "" if value is None or (isinstance(value, float) and pd.isna(value)) else str(value)
            )
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def _pairs(doc: Any, rows: list[tuple[str, str]]) -> None:
    _table(doc, pd.DataFrame(rows, columns=["Item", "Value"]))


def _exclusions(flow: ParticipantFlow) -> pd.DataFrame:
    rows = []
    for exit_ in flow.exits:
        if exit_.reasons:
            rows.extend(
                {"Step": exit_.label, "Criterion": name, "Rows excluded first": f"{n:,}"}
                for name, n in exit_.reasons
            )
        else:
            rows.append(
                {"Step": exit_.label, "Criterion": "", "Rows excluded first": f"{exit_.n:,}"}
            )
    return pd.DataFrame(rows, columns=["Step", "Criterion", "Rows excluded first"])


def build_dossier(inp: DossierInput) -> BytesIO:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading(inp.title, 0)
    meta = doc.add_paragraph()
    _small(
        meta,
        f"Generated on {inp.generated_at[:16].replace('T', ' ')}"
        + (f" by {inp.author}" if inp.author else "")
        + " with CohortStudio.",
    )
    note = doc.add_paragraph()
    _small(
        note,
        "This dossier supports the reporting of the participant flow (STROBE item 13) and of multiple imputation "
        "(Sterne et al. 2009). It documents what the application recorded; it does not certify compliance with a "
        "guideline, and it is not an electronically signed record.",
    )

    doc.add_heading("1. Provenance", level=1)
    _pairs(doc, inp.provenance)

    doc.add_heading("2. Analysis population and participant flow", level=1)
    doc.add_paragraph(f"Analysis population: {inp.population_label}.")
    if inp.flow_png:
        doc.add_picture(BytesIO(inp.flow_png), width=Cm(15))
        _small(doc.add_paragraph(), "Figure 1. Participant flow (STROBE item 13).")
    exclusions = _exclusions(inp.flow)
    if not exclusions.empty:
        doc.add_paragraph().add_run("Exclusions").bold = True
        _table(doc, exclusions)
    if inp.criteria is not None and not inp.criteria.empty:
        doc.add_paragraph().add_run("Criteria (ticked, in the order they are applied)").bold = True
        _table(doc, inp.criteria)
    for line in [
        *inp.flow.notes,
        *(f"The flow does not add up: {p}" for p in inp.flow.inconsistencies()),
    ]:
        _small(doc.add_paragraph(), line)

    doc.add_heading("3. Data handling", level=1)
    if inp.data_steps:
        for i, (name, text) in enumerate(inp.data_steps, start=1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. {name}: ").bold = True
            p.add_run(text)
    else:
        doc.add_paragraph("No change of the data was recorded after the load.")

    doc.add_heading("4. Missing data", level=1)
    if inp.missing is not None and not inp.missing.empty:
        doc.add_paragraph("Missing values in the analysis population:")
        _table(doc, inp.missing)
    else:
        doc.add_paragraph("No variable of the analysis population has missing values.")
    if inp.exact_imputed:
        total = sum(inp.exact_imputed.values())
        detail = ", ".join(f"{c}: {n:,}" for c, n in inp.exact_imputed.items())
        doc.add_paragraph(
            f"{total:,} values were imputed by an exact formula, the other variables determining them ({detail}); "
            "the formulas are listed under Data handling. Nothing was estimated, so the analyses treat them as "
            "observed."
        )
    if inp.single_imputed:
        total = sum(inp.single_imputed.values())
        detail = ", ".join(f"{c}: {n:,}" for c, n in inp.single_imputed.items())
        doc.add_paragraph(
            f"{total:,} values were filled by single imputation ({detail}). Analyses treat them as observed, so "
            "their standard errors are too small; the results that use them say so."
        )
    if inp.mi:
        doc.add_heading("Multiple imputation", level=2)
        doc.add_paragraph(inp.mi.get("methods_paragraph", ""))
        if inp.mi.get("settings"):
            _pairs(doc, inp.mi["settings"])
        if inp.mi.get("methods") is not None and not inp.mi["methods"].empty:
            doc.add_paragraph().add_run("Imputation method per variable").bold = True
            _table(doc, inp.mi["methods"])
        if inp.mi.get("status"):
            _small(doc.add_paragraph(), inp.mi["status"])
    else:
        doc.add_paragraph(
            "No multiple imputation was used: each analysis used the rows with every value it needs (complete "
            "cases); the number analysed is given with each result."
        )

    doc.add_heading("5. Results", level=1)
    if not inp.results:
        doc.add_paragraph("No result was added to the dossier.")
    figure = 1
    for result in inp.results:
        doc.add_heading(result.title, level=2)
        if result.settings:
            _small(doc.add_paragraph(), f"Settings: {result.settings}")
        _small(doc.add_paragraph(), result.provenance)
        if result.stale:
            p = doc.add_paragraph()
            run = p.add_run(
                f"Not current: {result.stale}. Kept at the user's request; run the analysis again to update it."
            )
            run.bold = True
        if result.single_imputed:
            detail = ", ".join(f"{c}: {n:,}" for c, n in result.single_imputed)
            _small(
                doc.add_paragraph(),
                f"Uses values filled by single imputation ({detail}): standard errors are too small.",
                italic=False,
            )
        for caption, frame in result.blocks.tables:
            if caption:
                doc.add_paragraph().add_run(caption).bold = True
            _table(doc, frame)
        for line in result.blocks.lines:
            doc.add_paragraph(line)
        if result.blocks.figure:
            figure += 1
            doc.add_picture(BytesIO(result.blocks.figure), width=Cm(13))
            _small(doc.add_paragraph(), f"Figure {figure}. {result.title}.")
        for warning in dict.fromkeys(result.blocks.warnings):
            _small(doc.add_paragraph(), f"Note: {warning}")
    if inp.omitted:
        _small(
            doc.add_paragraph(),
            "Left out because they were computed on other data or another population: "
            + "; ".join(inp.omitted)
            + ".",
        )

    doc.add_heading("6. References", level=1)
    for i, ref in enumerate(dict.fromkeys(inp.references), start=1):
        doc.add_paragraph(f"{i}. {ref}")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def exits_text(exit_: FlowExit) -> str:
    """One exclusion in words, for captions."""
    reasons = "; ".join(f"{name}: {n:,}" for name, n in exit_.reasons)
    return f"{exit_.label} (n = {exit_.n:,})" + (f": {reasons}" if reasons else "")


__all__ = ["METHOD_REFERENCES", "DossierInput", "DossierResult", "build_dossier", "exits_text"]
