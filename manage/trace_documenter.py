
import json
import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from core.lineage import verify_trace


class TraceDocumenter:
    def __init__(self, trace_source):
        """
        Initialize with either a dictionary trace or a path to a json trace file.
        """
        self.trace = {}
        if isinstance(trace_source, dict):
            self.trace = trace_source
        elif isinstance(trace_source, str) and os.path.exists(trace_source):
            try:
                with open(trace_source, 'r', encoding='utf-8') as f:  # JSON is UTF-8 (not the Windows code page)
                    self.trace = json.load(f)
            except Exception as e:
                import logging
                logging.getLogger(__name__).error("Error loading trace file: %s", e)
        
    def generate_report(self):
        """
        Generates a Word document report of the transformation trace.
        Returns the BytesIO object of the document.
        """
        from io import BytesIO
        
        document = Document()
        
        # Style Configuration
        style = document.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        
        # Title
        self._add_title(document)
        
        # Session Info
        self._add_session_info(document)
        
        # Steps
        document.add_heading('Transformation Steps', level=1)
        
        steps = self.trace.get("steps", [])
        if not steps:
            p = document.add_paragraph()
            p.add_run("No transformation steps recorded.").italic = True
        else:
            for i, step in enumerate(steps):
                self._add_step(document, i + 1, step)
                
        # Footer / Appendix?
        self._add_footer(document)
        
        # Save to buffer
        f = BytesIO()
        document.save(f)
        f.seek(0)
        return f

    def _add_title(self, doc):
        head = doc.add_heading('Data Transformation Report', 0)
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

    def _add_session_info(self, doc):
        doc.add_heading('Session Information', level=1)
        
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        self._add_row(table, "Session ID", self.trace.get("session_id", "N/A"))
        self._add_row(table, "Source Dataset", self.trace.get("source_dataset", "N/A"))
        self._add_row(table, "Step Count", str(len(self.trace.get("steps", []))))
        check = verify_trace(self.trace)
        integrity = {
            "verified": f"Verified: hash chain intact, head {(check.head_hash or '')[:16]}",
            "unverifiable": "Not verifiable: recorded before traces were chained",
        }.get(check.status, f"FAILED: {check.message}")
        self._add_row(table, "Integrity", integrity)
        parent = self.trace.get("parent")
        if parent:
            self._add_row(table, "Continues", ", ".join(f"{k}: {v}" for k, v in parent.items() if v is not None))

        doc.add_paragraph() # Spacer

    def _add_step(self, doc, index, step):
        # Step Header
        category = self._get_readable_function_name(step.get("function", "Unknown"))
        head = doc.add_heading(f"Step {index}: {category}", level=2)
        
        # Metadata line
        ts = step.get("timestamp", "")
        if ts:
            try:
                # Cleaning ISO format if needed
                ts = ts.replace("T", " ").split(".")[0]
            except:
                pass
        
        p_meta = doc.add_paragraph()
        run = p_meta.add_run(f"Executed at: {ts}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
        # Description
        desc = step.get("description", "")
        if desc:
            doc.add_paragraph(desc)

        # What the step did to the table (traces since schema 2)
        data = step.get("data") or {}
        if data.get("fp_out"):
            line = (
                f"Rows {data.get('rows_in', 0):,} -> {data.get('rows_out', 0):,} · "
                f"columns {data.get('cols_in', 0)} -> {data.get('cols_out', 0)} · "
                f"fingerprint {str(data.get('fp_in') or '-')[:8]} -> {str(data['fp_out'])[:8]}"
            )
            added = data.get("columns_added") or []
            if added:
                line += f" · added: {', '.join(added[:8])}{'...' if len(added) > 8 else ''}"
            run = doc.add_paragraph().add_run(line)
            run.font.size = Pt(9)
        if step.get("replayable") is False:
            run = doc.add_paragraph().add_run("This step cannot be recomputed by the replay engine.")
            run.italic = True
            run.font.size = Pt(9)

        population = step.get("population") or {}
        if population.get("criteria"):
            p = doc.add_paragraph()
            p.add_run(
                f"Analysis population: {population.get('n_included', 0):,} of {population.get('n_base', 0):,} rows"
            ).bold = True
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Light Shading'
            for cell, title in zip(table.rows[0].cells, ("Criteria family", "Rows in", "Excluded", "Rows out"), strict=True):
                cell.text = title
            for stage in population.get("stages") or []:
                cells = table.add_row().cells
                cells[0].text = f"{stage.get('family')} ({stage.get('kind')})"
                cells[1].text = f"{stage.get('n_in', 0):,}"
                cells[2].text = f"{stage.get('n_excluded', 0):,}"
                cells[3].text = f"{stage.get('n_out', 0):,}"
            
        # Multiple imputation: the methods paragraph (reporting items of Sterne et al. 2009)
        params = step.get("params", {})
        if step.get("function") == "mi_impute" and isinstance(params, dict) and params.get("spec"):
            self._add_mi_methods(doc, params)

        # Parameters Table
        if params:
            p = doc.add_paragraph()
            p.add_run("Parameters:").bold = True
            self._render_params(doc, params)
            
        doc.add_paragraph() # Spacer

    def _add_mi_methods(self, doc, params):
        """The methods paragraph of a multiple imputation, written from its recorded specification."""
        from types import SimpleNamespace

        from biostats.mi.report import methods_paragraph

        spec = params["spec"]
        binding = SimpleNamespace(
            spec=spec, methods=params.get("methods") or {}, columns=list(spec.get("columns") or [])
        )
        try:
            text = methods_paragraph(
                binding,
                n_rows=int(params.get("n_rows", 0)),
                incomplete=params.get("incomplete"),
                info={"engine": params.get("engine", "CohortStudio")},
            )
        except (KeyError, TypeError, ValueError):
            return
        p = doc.add_paragraph()
        p.add_run("Methods paragraph (reporting items of Sterne et al. 2009): ").bold = True
        p.add_run(text)

    def _render_params(self, doc, params, indent=0):
        if not params:
            return

        table = doc.add_table(rows=0, cols=2)
        table.style = 'Light Shading' # Lighter style for params
        
        for k, v in params.items():
            key_cell = table.add_row().cells
            key_cell[0].text = str(k)
            
            # Format value (long lists, such as kept row ids, are summarised)
            if isinstance(v, list):
                if len(v) > 10:
                    key_cell[1].text = f"{len(v)} items: {', '.join(str(x) for x in v[:5])}, ..."
                else:
                    key_cell[1].text = ", ".join([str(x) for x in v])
            elif isinstance(v, dict):
                 key_cell[1].text = json.dumps(v, indent=2)
            else:
                key_cell[1].text = str(v)

    def _add_row(self, table, label, value):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].text = str(value)

    def _get_readable_function_name(self, func_name):
        mapping = {
            "initial_load": "Dataset Loading",
            "enrichment": "Data Enrichment",
            "imputation": "Missing Value Imputation",
            "variable_transformation": "Variable Transformation",
            "variable_computation": "New Variable Computation",
            "outlier_handling": "Outlier Detection & Handling",
            "data_type_conversion": "Data Type Conversion",
            "missing_value_handling": "Manual Missing Value Handling",
            "apply_filter": "Filter Applied to the Dataset",
            "population_definition": "Analysis Population",
            "freeze_population": "Analysis Population Kept as the Dataset",
            "mi_impute": "Multiple Imputation",
            "targeted_imputation": "Targeted Imputation (Formula)",
            "multiple_imputation": "Multiple Imputation (Diagnostic, One Dataset Adopted)",
        }
        return mapping.get(func_name, func_name.replace("_", " ").title())

    def _add_footer(self, doc):
        p = doc.add_paragraph()
        run = p.add_run(
            "This report documents the steps recorded in the transformation trace. The trace is chained "
            "by hashes, so a later modification of a step is detectable (see Integrity); it is not an "
            "electronically signed record."
        )
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)
