"""
A value of the data never reaches what is shared (M5): strings planted in a malformed line and in a cell of an
imported file are absent from the trace, the trace report and the study dossier, and the trace names no absolute
path of the server. The malformed line is still shown on the Import page, where it helps the user.
"""

from __future__ import annotations

import io
import json
import os
from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import utils.data_paths as dp
from core.population import catalogue, configuration_for_sharing, configuration_problems

CANARY_LINE = "CANARY-LINE-9F2"
CANARY_CELL = "CANARY-CELL-4D1"


def _csv() -> bytes:
    rows = ["id,age,note"] + [f"{i},{40 + i % 30},routine" for i in range(1, 60)]
    rows.insert(5, f"5,47,{CANARY_LINE},extra")  # one field too many: skipped by the parser
    rows.append(f"60,79,{CANARY_CELL}")  # a value of one participant
    return "\n".join(rows).encode()


def _script() -> None:
    import streamlit as st

    from app_pages.study_dossier import build_dossier_bytes
    from ingest.loaders import read_table
    from ui.results import report_items
    from ui.state import get_state

    state = get_state()
    ss = st.session_state
    data = ss["csv"]
    result = read_table(data, "cohort.csv")
    state.load_parsed(
        result.df, filename="cohort.csv", data=data, options=result.options, name="cohort"
    )
    state.note("parse_report", result.report.to_dict(), "Parse report of the imported file")
    ss["examples"] = list(result.report.bad_line_examples)
    ss["trace"] = state.transformation_manager.get_trace()
    ss["dossier"] = build_dossier_bytes(
        state, report_items(), include_stale=False, title="Canary", author="tester"
    )


@pytest.fixture
def data_root(tmp_path: Path):
    previous = dp.DATA_ROOT
    dp.DATA_ROOT = str(tmp_path)
    os.environ.setdefault("CMP_IAM_PROVIDER", "local")
    yield tmp_path
    dp.DATA_ROOT = previous


def _docx_text(data: bytes | io.BytesIO) -> str:
    doc = Document(io.BytesIO(data) if isinstance(data, bytes) else data)
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_a_planted_value_never_reaches_a_shared_output(data_root: Path) -> None:
    from manage.trace_documenter import TraceDocumenter

    at = AppTest.from_function(_script, default_timeout=120)
    at.session_state["username"] = "canary_user"
    at.session_state["csv"] = _csv()
    at.run()
    assert not at.exception, [str(e.value)[:300] for e in at.exception]
    # the malformed line is shown on screen, and only there
    assert any(CANARY_LINE in example for example in at.session_state["examples"])

    trace = at.session_state["trace"]
    trace_text = json.dumps(trace, default=str)
    assert CANARY_LINE not in trace_text and CANARY_CELL not in trace_text
    for path in {str(data_root), str(data_root.resolve()), str(data_root).replace("\\", "\\\\")}:
        assert path not in trace_text  # the trace writes paths relative to the workspace
    load = trace["steps"][0]["params"]
    assert load["source_path"].startswith("traces/") and load["file_hash"].startswith("sha256:")
    assert {"cohortstudio", "python", "pandas"} <= set(load["environment"])

    trace_file = data_root / "trace.json"
    trace_file.write_text(trace_text, encoding="utf-8")
    report = _docx_text(TraceDocumenter(str(trace_file)).generate_report())
    assert CANARY_LINE not in report and CANARY_CELL not in report

    dossier = _docx_text(at.session_state["dossier"])
    assert CANARY_LINE not in dossier and CANARY_CELL not in dossier
    assert "SHA-256" in dossier  # the source file's checksum, with its algorithm


SENTINEL_ID = 73_196_451  # a row id planted in the rules that list rows


def test_a_configuration_for_sharing_lists_no_row() -> None:
    old = {"expression": "df['age'] > 80"}
    listed = {"row_ids": [[SENTINEL_ID, SENTINEL_ID]], "method": "zscore"}
    config = {
        "mask_families": {
            "outliers": {
                "operator": "any",
                "iforest": {"row_ids": [[3, 3], [SENTINEL_ID, SENTINEL_ID]], "columns": ["sbp"]},
                "frozen bounds": {"bounds": {"sbp": [80, 200]}},
            },
            "only_rows": {"operator": "any", "lof": {"row_ids": [[4, SENTINEL_ID]]}},
            "study_inclusion": {"operator": "all", "adult": {"expression": "df['age'] >= 18"}},
            # rules that combine others, one of which lists rows: at the first level, and at the second
            "combined": {
                "operator": "any",
                "two signs": {"threshold": 2, "members": {"listed": listed, "old": old}},
                "deeper": {
                    "threshold": 1,
                    "members": {"inner": {"threshold": 1, "members": {"listed": listed}}},
                },
                "high": {"expression": "df['sbp'] > 180"},
            },
            # a family counted with a threshold: without its listed rule, "at least 2 of" would count the others
            "warning_signs": {
                "operator": {"threshold": 2},
                "listed": listed,
                "sbp high": {"expression": "df['sbp'] > 180"},
                "dbp high": {"expression": "df['dbp'] > 110"},
            },
            # ids under keys that no kind holds, in a rule and in an operator
            "hand_written": {
                "operator": {"threshold": 1, "rows": [SENTINEL_ID]},
                "old": {**old, "flagged": [SENTINEL_ID]},
            },
            # ids where a single value is read: the copy never relies on the upload check, which refuses these
            "odd_values": {
                "operator": "any",
                "flag as a number": {"numeric": SENTINEL_ID, "column": "age", "lower_bound": 18},
                # texts where a number or a strategy is read: loaded, these fail its checks
                "bound in words": {
                    "numeric": True,
                    "column": "age",
                    "lower_bound": f"row {SENTINEL_ID}",
                },
                "strategy in words": {
                    "numeric": True,
                    "column": "age",
                    "lower_bound": 18,
                    "strategy": f"row {SENTINEL_ID}",
                },
                "bounds in words": {"bounds": {"age": [f"row {SENTINEL_ID}", None]}},
                "threshold in words": {"threshold": f"row {SENTINEL_ID}", "members": {"old": old}},
                "high": {"expression": "df['sbp'] > 180"},
            },
            "odd_operator": {"operator": {"threshold": [SENTINEL_ID]}, "old": old},
        },
        # and in the labels of a derived variable, which the upload check does not read
        "derived_variables": {
            "bmi": {
                "formula": "weight_kg / (height_cm / 100) ** 2",
                "kind": [SENTINEL_ID],
                "origin": {"rows": [SENTINEL_ID]},
            },
            "odd": {"formula": [SENTINEL_ID]},
        },
    }
    shared = configuration_for_sharing(config)
    text = json.dumps(shared.config)
    assert "row_ids" not in text and str(SENTINEL_ID) not in text
    families = shared.config["mask_families"]
    # a family left empty is dropped, and so is a threshold family that held a list, whole
    assert set(families) == {
        "outliers",
        "study_inclusion",
        "combined",
        "hand_written",
        "odd_values",
    }
    assert set(families["combined"]) == {"operator", "high"}
    assert families["hand_written"] == {"operator": {"threshold": 1}, "old": old}
    assert families["outliers"] == {
        "operator": "any",
        "frozen bounds": {"bounds": {"sbp": [80, 200]}},
    }
    assert shared.config["derived_variables"] == {
        "bmi": {"formula": "weight_kg / (height_cm / 100) ** 2"}
    }
    assert shared.with_rows == (
        "outliers: iforest",
        "only_rows: lof",
        "combined: two signs",
        "combined: deeper",
    )
    assert shared.families == ("warning_signs",)
    assert shared.unreadable == (
        "odd_values: flag as a number",
        "odd_values: bound in words",
        "odd_values: strategy in words",
        "odd_values: bounds in words",
        "odd_values: threshold in words",
        "odd_operator",
        "derived_variables: odd",
    )
    # the derived variable "bmi" is shared without the kind and origin it cannot be read with
    assert shared.trimmed == (
        "hand_written: old",
        "hand_written: the operator",
        "derived_variables: bmi",
    )
    assert configuration_problems(shared.config) == []  # the file loads
    assert str(SENTINEL_ID) in json.dumps(config)  # the session's configuration is unchanged


def test_a_study_shared_keeps_the_fingerprints_of_its_criteria() -> None:
    """The file for sharing keeps every key the app writes in a rule: its criteria are the study's, fingerprints
    included."""
    path = Path(__file__).resolve().parents[2] / "assets" / "tutorial" / "tutorial_config.json"
    tutorial = json.loads(path.read_text(encoding="utf-8"))
    shared = configuration_for_sharing(tutorial)
    assert shared.with_rows == shared.families == shared.unreadable == shared.trimmed == ()
    assert [c.fingerprint for c in catalogue(shared.config)] == [
        c.fingerprint for c in catalogue(tutorial)
    ]


def _download_script() -> None:
    """Validate's own buttons: the file each one gives, by its name."""
    from typing import Any

    import pandas as pd
    import streamlit as st

    from app_pages.data_monitoring import display_results
    from ui.state import get_state

    get_state().load_tutorial_cohort()
    listed = {"row_ids": [[st.session_state["id"], st.session_state["id"]]], "method": "zscore"}
    st.session_state.config["mask_families"]["outliers"] = {"operator": "any", "flagged": listed}
    files: dict[str, Any] = {}
    button = st.download_button

    def keep(*args: Any, **kwargs: Any) -> Any:
        files[str(kwargs.get("file_name"))] = kwargs.get("data")
        return button(*args, **kwargs)

    st.download_button = keep  # type: ignore[assignment]
    try:
        frame = pd.DataFrame({"x": [1]})
        display_results(frame, None, None, "anomaly_data", st)  # pyright: ignore[reportArgumentType]
    finally:
        st.download_button = button  # type: ignore[assignment]
    st.session_state["files"] = files


def test_validate_gives_the_file_for_sharing_under_its_own_name(data_root: Path) -> None:
    """The page's buttons: config.json is whole, to load again with this dataset; config_for_sharing.json holds no
    row of it."""
    at = AppTest.from_function(_download_script, default_timeout=120)
    at.session_state["username"] = "canary_user"
    at.session_state["id"] = SENTINEL_ID
    at.run()
    assert not at.exception, [str(e.value)[:300] for e in at.exception]
    files = at.session_state["files"]
    assert str(SENTINEL_ID) in files["config.json"]
    assert str(SENTINEL_ID) not in files["config_for_sharing.json"]
    assert str(data_root).replace("\\", "\\\\") not in files["config_for_sharing.json"]


def _sharing_script() -> None:
    """The rules, the enrichment and the derived variable written as the pages write them, then the file for sharing."""
    from pathlib import Path

    import streamlit as st

    from app_pages.data_enrichment import save_enrichment_config
    from app_pages.data_monitoring import _exclude_outliers, left_out_text
    from core.derived import CONFIG_KEY
    from core.population import configuration_for_sharing
    from ui.state import get_state

    state = get_state()
    ss = st.session_state
    state.load_tutorial_cohort()  # its configuration names this installation's folders
    detection = {"method": "zscore", "params": {"threshold": 3.0}, "columns": ["sbp_baseline"]}
    _exclude_outliers({**detection, "flagged": ss["ids"]})
    upload = Path(ss["upload"])
    save_enrichment_config(
        "labs",
        "merge",
        input_file_folder=str(upload.parent),
        input_file=str(upload),
        identifier="id",
    )
    ss.config[CONFIG_KEY] = {
        "bmi": {"formula": "weight_kg / (height_cm / 100) ** 2", "note": ss["note"]}
    }
    shared = configuration_for_sharing(ss.config)
    ss["shared"] = shared.config
    ss["caption"] = left_out_text(shared)


def test_a_file_for_sharing_holds_no_row_path_or_note(data_root: Path) -> None:
    at = AppTest.from_function(_sharing_script, default_timeout=120)
    at.session_state["username"] = "canary_user"
    at.session_state["ids"] = [SENTINEL_ID, SENTINEL_ID + 1]
    at.session_state["upload"] = str(
        data_root / "users" / "canary_user" / "uploads" / "labs_site_B.xlsx"
    )
    at.session_state["note"] = CANARY_NOTE
    at.run()
    assert not at.exception, [str(e.value)[:300] for e in at.exception]
    shared = json.dumps(at.session_state["shared"], default=str)  # as the page writes the file
    for planted in (
        str(SENTINEL_ID),
        str(SENTINEL_ID + 1),
        CANARY_NOTE,
        "labs_site_B",
        "canary_user",
    ):
        assert planted not in shared, planted
    for path in {str(data_root), str(data_root.resolve()), str(data_root).replace("\\", "\\\\")}:
        assert path not in shared
    config = json.loads(shared)
    assert set(config) == {"description", "mask_families", "derived_variables"}
    assert config["derived_variables"] == {"bmi": {"formula": "weight_kg / (height_cm / 100) ** 2"}}
    assert configuration_problems(config) == []  # the file loads
    caption = at.session_state["caption"]
    assert "zscore outliers (sbp_baseline)" in caption
    assert (
        "the notes of the derived variables" in caption
        and "this installation's settings" in caption
    )


CANARY_LEVEL = "CANARY-LEVEL-7A3"
CANARY_NOTE = "CANARY-NOTE-2B8"


def _fill_script() -> None:
    import numpy as np
    import pandas as pd
    import streamlit as st

    from app_pages.study_dossier import build_dossier_bytes
    from core.derived import DerivedVariable
    from core.state import DatasetRef
    from ui.results import report_items
    from ui.state import get_state

    state = get_state()
    ss = st.session_state
    rng = np.random.default_rng(3)
    n = 200
    sex = rng.choice(["F", "M"], n).astype(object)
    x = np.round(rng.normal(50, 10, n), 0)
    frame = pd.DataFrame(
        {"x": x, "sex": sex, "y": np.round(x * np.where(sex == "F", 0.85, 1.0), 1)}
    )
    frame.loc[0:9, "y"] = np.nan
    frame.loc[[0, 50], "sex"] = ss["level"]  # held by one row compared, and one to fill
    frame.loc[50, "y"] = frame.loc[50, "x"]
    state.set_dataset(frame, DatasetRef(name="d", origin="upload", file_hash="h"))
    state.set_derived_variables(
        [DerivedVariable("y", "x * np.where(sex == 'F', 0.85, 1.0)", note=ss["note"])]
    )
    ss["filled"] = state.fill_exactly()
    ss["applied"] = state.commit_staging()
    ss["trace"] = state.transformation_manager.get_trace()
    ss["dossier"] = build_dossier_bytes(
        state, report_items(), include_stale=False, title="Canary", author="tester"
    )


def test_a_rare_level_or_a_note_never_reaches_a_shared_output_through_a_fill(
    data_root: Path,
) -> None:
    from manage.trace_documenter import TraceDocumenter

    at = AppTest.from_function(_fill_script, default_timeout=120)
    at.session_state["username"] = "canary_user"
    at.session_state["level"], at.session_state["note"] = CANARY_LEVEL, CANARY_NOTE
    at.run()
    assert not at.exception, [str(e.value)[:300] for e in at.exception]
    assert at.session_state["filled"] == 9  # the row with the rare level is not filled
    trace = at.session_state["trace"]
    trace_text = json.dumps(trace, default=str)
    assert "targeted_imputation" in trace_text and '"only_levels"' in trace_text
    assert CANARY_LEVEL not in trace_text and CANARY_NOTE not in trace_text
    trace_file = data_root / "trace.json"
    trace_file.write_text(trace_text, encoding="utf-8")
    report = _docx_text(TraceDocumenter(str(trace_file)).generate_report())
    assert CANARY_LEVEL not in report and CANARY_NOTE not in report
    dossier = _docx_text(at.session_state["dossier"])
    assert CANARY_LEVEL not in dossier and CANARY_NOTE not in dossier
