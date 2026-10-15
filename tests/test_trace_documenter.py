import json
from pathlib import Path

from docx import Document

from manage.trace_documenter import TraceDocumenter


def test_trace_documenter(tmp_path: Path):
    """Test generating docx reports from trace files."""
    trace_data = {
        "session_id": "test_session",
        "source_dataset": "test_source.csv",
        "steps": [
            {
                "timestamp": "2026-01-14T12:00:00",
                "function": "enrichment",
                "description": "Test Enrichment",
                "params": {"file": "external.csv", "cols": ["A", "B"]},
            }
        ],
    }
    trace_path = tmp_path / "test_trace.json"
    trace_path.write_text(json.dumps(trace_data), encoding="utf-8")

    doc_generator = TraceDocumenter(str(trace_path))
    buffer = doc_generator.generate_report()

    assert buffer is not None
    assert buffer.getbuffer().nbytes > 0

    # Verify generated DOCX structure
    doc = Document(buffer)
    text = "\n".join([p.text for p in doc.paragraphs])
    assert "Data Transformation Report" in text
    assert "Session Information" in text


def test_trace_file_is_read_as_utf8(tmp_path: Path):
    """A trace file in UTF-8 (characters not escaped) is read as UTF-8, not in the Windows code page."""
    trace_data = {
        "session_id": "s",
        "source_dataset": "cohorte.csv",
        "steps": [
            {
                "timestamp": "2026-10-14T12:00:00",
                "function": "rename_columns",
                "description": "Renommé « poids » en poids_kg",
                "params": {},
            }
        ],
    }
    trace_path = tmp_path / "trace.json"
    trace_path.write_text(json.dumps(trace_data, ensure_ascii=False), encoding="utf-8")

    doc = Document(TraceDocumenter(str(trace_path)).generate_report())
    assert "Renommé « poids » en poids_kg" in "\n".join(p.text for p in doc.paragraphs)


def test_a_multiple_imputation_step_gets_its_methods_paragraph(tmp_path: Path):
    """The report writes the methods paragraph of an mi_impute step from its recorded specification."""
    spec = {
        "columns": ["age", "hba1c_pct", "death"],
        "m": 20,
        "maxit": 10,
        "seed": 2026,
        "donors": 5,
        "by": None,
        "survival": ["fu_days", "death"],
    }
    trace_data = {
        "session_id": "s",
        "source_dataset": "cohort.csv",
        "steps": [
            {
                "timestamp": "2026-10-15T12:00:00",
                "function": "mi_impute",
                "description": "Multiple imputation, m = 20: hba1c_pct (pmm)",
                "params": {
                    "spec": spec,
                    "methods": {"hba1c_pct": "pmm"},
                    "m": 20,
                    "n_rows": 928,
                    "n_imputed": 4200,
                    "incomplete": 0.226,
                    "engine": "cohortstudio-fcs 1",
                },
            }
        ],
    }
    trace_path = tmp_path / "trace.json"
    trace_path.write_text(json.dumps(trace_data), encoding="utf-8")

    text = "\n".join(p.text for p in Document(TraceDocumenter(str(trace_path)).generate_report()).paragraphs)
    assert "Methods paragraph (reporting items of Sterne et al. 2009)" in text
    assert "Of the 928 participants of the analysis population, 22.6% had at least one missing value" in text
    assert "We created 20 imputed datasets, each after 10 iterations (seed 2026; cohortstudio-fcs 1" in text
    assert "Nelson-Aalen" in text  # the planned Cox model
